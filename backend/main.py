import os
import hashlib
import csv
import io
import json
import logging
import re
import unicodedata
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from fastapi import (
    Body,
    Depends,
    File,
    FastAPI,
    Form,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import inspect, text
from sqlalchemy.orm import Session

import crud
import models
import schemas
from database import Base, SessionLocal, engine
from security import authenticate_user, create_access_token, get_current_admin
from services.blockchain_service import (
    record_external_request_rejection_on_blockchain,
    record_student_rejection_on_blockchain,
    sync_student_to_blockchain,
    verify_student_on_blockchain,
)
from services.encryption_service import EncryptionError, decrypt_value
from services.explorer_service import (
    get_block_detail,
    get_transaction_detail,
    list_blocks_page,
    list_transactions_page,
    list_recent_blocks,
    list_recent_transactions,
    search_explorer,
)
from services.email_service import (
    notify_admins_new_certificate_request,
    notify_requester_certificate_approved,
    notify_requester_certificate_rejected,
    send_certificate_request_otp_email,
)
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

Base.metadata.create_all(bind=engine)


def ensure_student_hash_version_column():
    inspector = inspect(engine)
    columns = {
        column["name"]
        for column in inspector.get_columns("students")
    }
    if "metadata_hash_version" in columns:
        return

    with engine.begin() as connection:
        connection.execute(text(
            "ALTER TABLE students ADD COLUMN metadata_hash_version "
            "VARCHAR(10) NOT NULL DEFAULT 'v1'"
        ))


ensure_student_hash_version_column()


# EXTERNAL_REQUEST_FABRIC_DECISION_V1
def ensure_external_request_blockchain_columns():
    inspector = inspect(engine)
    columns = {
        column["name"]
        for column in inspector.get_columns("external_requests")
    }
    statements = {
        "source_data_hash": (
            "ALTER TABLE external_requests ADD COLUMN "
            "source_data_hash VARCHAR(64) NULL"
        ),
        "blockchain_tx_id": (
            "ALTER TABLE external_requests ADD COLUMN "
            "blockchain_tx_id VARCHAR(100) NULL"
        ),
        "blockchain_action": (
            "ALTER TABLE external_requests ADD COLUMN "
            "blockchain_action VARCHAR(50) NULL"
        ),
    }

    with engine.begin() as connection:
        for column_name, statement in statements.items():
            if column_name not in columns:
                connection.execute(text(statement))


ensure_external_request_blockchain_columns()

app = FastAPI(
    title="Graduation Verification API",
    version="0.1.0",
)

cors_origins = [
    origin.strip().rstrip("/")
    for origin in os.getenv("CORS_ORIGINS", "").split(",")
    if origin.strip()
]

if not cors_origins:
    raise RuntimeError("CORS_ORIGINS is not configured")

app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()


def now_utc():
    return datetime.now(timezone.utc)


# EXTERNAL_REQUEST_FABRIC_DECISION_V1
def calculate_external_request_source_hash(external_request) -> str:
    try:
        imported_data = json.loads(external_request.import_data or "{}")
    except (TypeError, json.JSONDecodeError):
        imported_data = {"raw": str(external_request.import_data or "")}

    payload = {
        "hash_version": "external-source-v1",
        "source_file_name": external_request.source_file_name,
        "source_row_number": external_request.source_row_number,
        "student_id": external_request.student_id,
        "import_data": imported_data,
    }
    canonical_json = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(canonical_json.encode("utf-8")).hexdigest()


def classification_from_gpa(gpa) -> str | None:
    if gpa is None:
        return None

    gpa_value = Decimal(str(gpa))

    if gpa_value >= Decimal("3.60"):
        return "Excellent"

    if gpa_value >= Decimal("3.20"):
        return "Very Good"

    if gpa_value >= Decimal("2.50"):
        return "Good"

    if gpa_value >= Decimal("2.00"):
        return "Average"

    return "Weak"


MAX_EXTERNAL_IMPORT_BYTES = 5 * 1024 * 1024
MAX_EXTERNAL_IMPORT_ROWS = 1000

EXTERNAL_STUDENT_IMPORT_FIELDS = {
    "student_id": {
        "label": "Mã số sinh viên",
        "aliases": {"student_id", "mssv", "ma_so_sinh_vien", "ma_sinh_vien"},
    },
    "full_name": {
        "label": "Họ và tên",
        "aliases": {"full_name", "ho_ten", "ho_va_ten", "name"},
    },
    "date_of_birth": {
        "label": "Ngày sinh",
        "aliases": {"date_of_birth", "ngay_sinh", "birth_date"},
    },
    "citizen_id": {
        "label": "Số CCCD/CMND",
        "aliases": {"citizen_id", "cccd", "cmnd", "so_cccd", "so_cmnd"},
    },
    "email": {
        "label": "Email",
        "aliases": {"email", "student_email", "email_sinh_vien"},
    },
    "institution_code": {
        "label": "Mã trường",
        "aliases": {"institution_code", "ma_truong", "school_code"},
    },
    "institution_name": {
        "label": "Tên trường",
        "aliases": {"institution_name", "ten_truong", "truong", "school_name"},
    },
    "faculty_name": {
        "label": "Khoa",
        "aliases": {"faculty_name", "faculty", "khoa"},
    },
    "major": {
        "label": "Chuyên ngành",
        "aliases": {"major", "chuyen_nganh", "nganh_hoc"},
    },
    "training_mode": {
        "label": "Hình thức đào tạo",
        "aliases": {"training_mode", "hinh_thuc_dao_tao", "he_dao_tao"},
    },
    "degree_type": {
        "label": "Loại bằng",
        "aliases": {"degree_type", "loai_bang"},
    },
    "entrance_year": {
        "label": "Khóa nhập học",
        "aliases": {"entrance_year", "khoa_nhap_hoc", "nam_nhap_hoc"},
    },
    "graduation_status": {
        "label": "Trạng thái tốt nghiệp",
        "aliases": {"graduation_status", "trang_thai_tot_nghiep", "trang_thai"},
    },
    "graduation_date": {
        "label": "Ngày tốt nghiệp",
        "aliases": {"graduation_date", "ngay_tot_nghiep"},
    },
    "graduation_year": {
        "label": "Năm tốt nghiệp",
        "aliases": {"graduation_year", "nam_tot_nghiep"},
    },
    "gpa": {
        "label": "GPA",
        "aliases": {"gpa", "diem_gpa"},
    },
    "classification": {
        "label": "Xếp loại",
        "aliases": {"classification", "xep_loai"},
    },
    "total_credits": {
        "label": "Tổng tín chỉ",
        "aliases": {"total_credits", "tong_tin_chi"},
    },
    "expected_graduation_date": {
        "label": "Ngày tốt nghiệp dự kiến",
        "aliases": {"expected_graduation_date", "ngay_tot_nghiep_du_kien"},
    },
    "academic_status": {
        "label": "Tình trạng học tập",
        "aliases": {"academic_status", "tinh_trang_hoc_tap"},
    },
}

EXTERNAL_STUDENT_REQUIRED_FIELDS = {
    "student_id",
    "full_name",
    "date_of_birth",
    "citizen_id",
    "email",
    "institution_name",
    "faculty_name",
    "major",
    "training_mode",
    "degree_type",
    "entrance_year",
    "graduation_status",
}


def normalize_external_text(value) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        character
        for character in text
        if not unicodedata.combining(character)
    )
    return text.replace("đ", "d").replace("Đ", "D").lower()


FACULTY_NAME_ALIASES = {
    "information technology": "Công nghệ thông tin",
    "faculty of information technology": "Công nghệ thông tin",
    "school of information technology": "Công nghệ thông tin",
    "it": "Công nghệ thông tin",
    "cong nghe thong tin": "Công nghệ thông tin",
    "communication technology": "Công nghệ truyền thông",
    "communications technology": "Công nghệ truyền thông",
    "faculty of communication technology": "Công nghệ truyền thông",
    "cong nghe truyen thong": "Công nghệ truyền thông",
    "language": "Ngôn ngữ",
    "languages": "Ngôn ngữ",
    "foreign languages": "Ngôn ngữ",
    "ngon ngu": "Ngôn ngữ",
    "law": "Luật",
    "faculty of law": "Luật",
    "luat": "Luật",
    "business administration": "Quản trị kinh doanh",
    "business management": "Quản trị kinh doanh",
    "faculty of business administration": "Quản trị kinh doanh",
    "quan tri kinh doanh": "Quản trị kinh doanh",
    "computer science": "Khoa học máy tính",
    "faculty of computer science": "Khoa học máy tính",
    "khoa hoc may tinh": "Khoa học máy tính",
}


def normalize_faculty_name(value) -> str:
    original_value = str(value or "").strip()
    if not original_value:
        return original_value

    normalized_value = normalize_external_text(original_value)
    return FACULTY_NAME_ALIASES.get(normalized_value, original_value)


def normalize_external_header(value) -> str:
    normalized = normalize_external_text(value)
    return re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")


def clean_external_cell(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    cleaned = str(value).strip()
    return cleaned or None


def parse_external_date(value: str | None) -> date | None:
    if not value:
        return None

    cleaned = str(value).strip()
    for date_format in (
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(cleaned, date_format).date()
        except ValueError:
            continue
    return None


def parse_external_integer(value: str | None) -> int | None:
    if value is None:
        return None
    try:
        number = Decimal(str(value).replace(",", "."))
    except Exception:
        return None
    if number != number.to_integral_value():
        return None
    return int(number)


def parse_external_decimal(value: str | None) -> Decimal | None:
    if value is None:
        return None
    try:
        return Decimal(str(value).replace(",", "."))
    except Exception:
        return None


def normalize_external_status(value: str | None) -> str | None:
    normalized = normalize_external_header(value)
    if normalized in {"graduated", "da_tot_nghiep"}:
        return "GRADUATED"
    if normalized in {"not_graduated", "chua_tot_nghiep"}:
        return "NOT_GRADUATED"
    return None


def normalize_external_classification(value: str | None) -> str | None:
    normalized = normalize_external_text(value)
    values = {
        "excellent": "Excellent",
        "xuat sac": "Excellent",
        "very good": "Very Good",
        "gioi": "Very Good",
        "good": "Good",
        "kha": "Good",
        "average": "Average",
        "trung binh": "Average",
        "weak": "Weak",
        "yeu": "Weak",
    }
    return values.get(normalized)


def is_fpt_institution(imported_data: dict) -> bool:
    institution_code = normalize_external_header(
        imported_data.get("institution_code")
    )
    institution_name = normalize_external_text(
        imported_data.get("institution_name")
    )
    return institution_code in {"fpt", "fptu"} or "fpt" in institution_name


def parse_external_student_file(
    filename: str,
    file_content: bytes,
) -> list[dict]:
    extension = Path(filename).suffix.lower()

    if extension == ".csv":
        try:
            text = file_content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=422,
                detail="Tệp CSV phải sử dụng mã hóa UTF-8.",
            )

        try:
            dialect = csv.Sniffer().sniff(
                text[:4096],
                delimiters=",;|\t",
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = reader.fieldnames or []
        rows = [
            (row_number, row)
            for row_number, row in enumerate(reader, start=2)
        ]
    elif extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(
                status_code=500,
                detail="Backend chưa cài thư viện openpyxl.",
            ) from exc

        workbook = load_workbook(
            io.BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        value_rows = worksheet.iter_rows(values_only=True)
        headers = list(next(value_rows, []) or [])
        rows = [
            (
                row_number,
                {
                    headers[index]: (
                        values[index]
                        if index < len(values)
                        else None
                    )
                    for index in range(len(headers))
                },
            )
            for row_number, values in enumerate(value_rows, start=2)
        ]
        workbook.close()
    else:
        raise HTTPException(
            status_code=400,
            detail="Chỉ chấp nhận tệp CSV hoặc XLSX.",
        )

    if not headers:
        raise HTTPException(
            status_code=422,
            detail="Tệp không có dòng tiêu đề.",
        )
    if len(rows) > MAX_EXTERNAL_IMPORT_ROWS:
        raise HTTPException(
            status_code=400,
            detail=f"Tệp chỉ được chứa tối đa {MAX_EXTERNAL_IMPORT_ROWS} dòng.",
        )

    normalized_headers = {
        normalize_external_header(header): header
        for header in headers
        if header is not None
    }
    column_map = {
        field: next(
            (
                normalized_headers[alias]
                for alias in config["aliases"]
                if alias in normalized_headers
            ),
            None,
        )
        for field, config in EXTERNAL_STUDENT_IMPORT_FIELDS.items()
    }

    missing_columns = [
        EXTERNAL_STUDENT_IMPORT_FIELDS[field]["label"]
        for field in EXTERNAL_STUDENT_REQUIRED_FIELDS
        if column_map[field] is None
    ]
    if missing_columns:
        raise HTTPException(
            status_code=422,
            detail=(
                "Tệp còn thiếu các cột bắt buộc: "
                + ", ".join(sorted(missing_columns))
                + "."
            ),
        )

    parsed_rows = []
    for row_number, row in rows:
        imported_data = {
            field: clean_external_cell(row.get(header)) if header else None
            for field, header in column_map.items()
        }
        imported_data["student_id"] = str(
            imported_data.get("student_id") or ""
        ).strip().upper()

        if not any(imported_data.values()):
            continue

        parsed_rows.append(
            {
                "row_number": row_number,
                "data": imported_data,
            }
        )
    return parsed_rows


def prepare_external_student(
    db: Session,
    imported_data: dict,
) -> tuple[schemas.StudentCreate | None, str, str]:
    missing_fields = [
        EXTERNAL_STUDENT_IMPORT_FIELDS[field]["label"]
        for field in EXTERNAL_STUDENT_REQUIRED_FIELDS
        if not str(imported_data.get(field) or "").strip()
    ]
    if missing_fields:
        return (
            None,
            "MISSING_REQUIRED_DATA",
            "Thiếu dữ liệu bắt buộc: " + ", ".join(sorted(missing_fields)) + ".",
        )

    student_id = str(imported_data["student_id"]).strip().upper()
    citizen_id = str(imported_data["citizen_id"]).strip()
    email = str(imported_data["email"]).strip().lower()
    date_of_birth = parse_external_date(imported_data.get("date_of_birth"))
    entrance_year = parse_external_integer(imported_data.get("entrance_year"))
    graduation_status = normalize_external_status(
        imported_data.get("graduation_status")
    )
    graduation_date = parse_external_date(imported_data.get("graduation_date"))
    graduation_year = parse_external_integer(imported_data.get("graduation_year"))
    gpa = parse_external_decimal(imported_data.get("gpa"))
    total_credits = parse_external_integer(imported_data.get("total_credits"))
    expected_graduation_date = parse_external_date(
        imported_data.get("expected_graduation_date")
    )

    errors = []
    if not re.fullmatch(r"[A-Z]{2}\d{6}", student_id):
        errors.append("Mã số sinh viên phải gồm 2 chữ cái in hoa và 6 chữ số")
    if not re.fullmatch(r"\d{12}", citizen_id):
        errors.append("Số CCCD/CMND phải gồm đúng 12 chữ số")
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        errors.append("Email không đúng định dạng")
    if date_of_birth is None:
        errors.append("Ngày sinh không hợp lệ")
    if entrance_year is None:
        errors.append("Khóa nhập học không hợp lệ")
    if graduation_status is None:
        errors.append("Trạng thái tốt nghiệp không hợp lệ")
    if total_credits is not None and total_credits < 0:
        errors.append("Tổng tín chỉ không hợp lệ")

    if graduation_status == "NOT_GRADUATED":
        if gpa is not None:
            errors.append(
                "Hồ sơ không thống nhất: sinh viên chưa tốt nghiệp "
                "nhưng đã có điểm GPA"
            )

        # Hồ sơ chưa tốt nghiệp vẫn được tiếp nhận vào Quản lý sinh viên.
        # Các thông tin chỉ có sau khi tốt nghiệp không được lấy từ tệp nguồn.
        graduation_date = None
        graduation_year = None
        gpa = None

    if graduation_status == "GRADUATED":
        if graduation_date is None:
            errors.append("Sinh viên đã tốt nghiệp phải có ngày tốt nghiệp")
        if graduation_year is None:
            errors.append("Sinh viên đã tốt nghiệp phải có năm tốt nghiệp")
        if gpa is None:
            errors.append(
                "Không thể duyệt hồ sơ vì sinh viên chưa có điểm GPA"
            )
        elif gpa < Decimal("0") or gpa > Decimal("4"):
            errors.append("GPA phải từ 0.00 đến 4.00")

    if errors:
        return None, "INVALID_SOURCE_DATA", "; ".join(errors) + "."

    if not is_fpt_institution(imported_data):
        return (
            None,
            "NOT_FPT_INSTITUTION",
            "Hồ sơ không thuộc Trường Đại học FPT.",
        )

    existing_student = crud.get_student(db, student_id)
    if existing_student:
        return (
            None,
            "DUPLICATE_STUDENT_ID",
            "Mã số sinh viên đã có trong Quản lý Sinh viên.",
        )

    existing_email = crud.get_student_by_email(db, email)
    if existing_email:
        return (
            None,
            "DUPLICATE_EMAIL",
            "Email đã được sử dụng bởi sinh viên khác.",
        )

    citizen_id_hash = crud.calculate_citizen_id_hash(citizen_id)
    existing_citizen = crud.get_student_by_citizen_id_hash(
        db,
        citizen_id_hash,
    )
    if existing_citizen:
        return (
            None,
            "DUPLICATE_CITIZEN_ID",
            "Số CCCD/CMND đã được sử dụng bởi sinh viên khác.",
        )

    calculated_classification = classification_from_gpa(gpa)
    imported_classification = normalize_external_classification(
        imported_data.get("classification")
    )
    if (
        imported_classification
        and calculated_classification
        and imported_classification != calculated_classification
    ):
        return (
            None,
            "INVALID_SOURCE_DATA",
            "Xếp loại trong tệp không phù hợp với GPA.",
        )

    try:
        student_data = schemas.StudentCreate(
            student_id=student_id,
            full_name=str(imported_data["full_name"]).strip(),
            date_of_birth=date_of_birth,
            citizen_id=citizen_id,
            email=email,
            institution_code=(
                str(imported_data.get("institution_code") or "FPTU")
                .strip()
                .upper()
            ),
            institution_name=str(imported_data["institution_name"]).strip(),
            faculty_name=normalize_faculty_name(
                imported_data.get("faculty_name")
            ),
            major=str(imported_data["major"]).strip(),
            training_mode=str(imported_data["training_mode"]).strip(),
            degree_id=None,
            degree_type=str(imported_data["degree_type"]).strip(),
            entrance_year=entrance_year,
            graduation_status=graduation_status,
            graduation_date=graduation_date,
            graduation_year=graduation_year,
            classification=calculated_classification,
            gpa=gpa,
            total_credits=total_credits,
            expected_graduation_date=expected_graduation_date,
            academic_status=(
                str(imported_data.get("academic_status") or "").strip()
                or None
            ),
        )
    except Exception as exc:
        logger.info("Invalid external student payload %s: %s", student_id, exc)
        return (
            None,
            "INVALID_SOURCE_DATA",
            "Dữ liệu hồ sơ không đáp ứng quy tắc của hệ thống.",
        )

    if graduation_status == "NOT_GRADUATED":
        verification_message = (
            "Hồ sơ đủ thông tin cơ bản và sẵn sàng để duyệt. "
            "Sinh viên sẽ được thêm vào Quản lý sinh viên với trạng thái "
            "Chưa tốt nghiệp."
        )
    else:
        verification_message = (
            "Hồ sơ đầy đủ và sẵn sàng để admin xem xét."
        )

    return student_data, "READY_FOR_REVIEW", verification_message


def raise_duplicate_field_errors(errors: list[dict]) -> None:
    if errors:
        raise HTTPException(
            status_code=409,
            detail=errors,
        )


def is_same_student_id(left: str | None, right: str | None) -> bool:
    return str(left or "").strip().upper() == str(right or "").strip().upper()


def validate_unique_student_fields(
    db: Session,
    student_data,
    current_student_id: str | None = None,
):
    errors = []
    student_id = getattr(student_data, "student_id", None)

    if student_id:
        existing_student = crud.get_student_including_deleted(db, student_id)

        if existing_student and not is_same_student_id(
            existing_student.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "student_id"],
                    "msg": (
                        "Mã số sinh viên đang nằm trong danh sách đã xóa. "
                        "Vui lòng khôi phục sinh viên thay vì thêm mới."
                        if existing_student.is_deleted
                        else "Mã số sinh viên đã tồn tại trong hệ thống."
                    ),
                    "type": "value_error.duplicate",
                }
            )

    degree_id = getattr(student_data, "degree_id", None)

    if degree_id:
        existing_degree = crud.get_student_by_degree_id(
            db,
            degree_id,
        )

        if existing_degree and not is_same_student_id(
            existing_degree.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "degree_id"],
                    "msg": "Degree already exists",
                    "type": "value_error.duplicate",
                }
            )

    email = getattr(student_data, "email", None)

    if email:
        existing_email = crud.get_student_by_email(db, email)

        if existing_email and not is_same_student_id(
            existing_email.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "email"],
                    "msg": "Email already exists",
                    "type": "value_error.duplicate",
                }
            )

    citizen_id = getattr(student_data, "citizen_id", None)

    if citizen_id:
        citizen_id_hash = crud.calculate_citizen_id_hash(citizen_id)

        existing_citizen = crud.get_student_by_citizen_id_hash(
            db,
            citizen_id_hash,
        )

        if existing_citizen and not is_same_student_id(
            existing_citizen.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "citizen_id"],
                    "msg": "Citizen ID already exists",
                    "type": "value_error.duplicate",
                }
            )

    raise_duplicate_field_errors(errors)

@app.get("/")
def root():
    return {"message": "Graduation Verification API is running"}


@app.post("/api/auth/login", response_model=schemas.TokenResponse)
def login(
    login_data: schemas.LoginRequest,
    db: Session = Depends(get_db),
):
    user = authenticate_user(
        db=db,
        username=login_data.username,
        password=login_data.password,
    )

    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid username or password",
        )

    user.last_login = now_utc()
    db.commit()

    access_token = create_access_token(
        data={
            "sub": user.username,
            "role": user.role,
        }
    )

    return {
        "access_token": access_token,
        "token_type": "bearer",
    }


@app.get("/api/admin/stats", response_model=schemas.StatsResponse)
def admin_stats(
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    return crud.get_stats(db)

@app.get(
    "/api/admin/audit-logs",
    response_model=schemas.AuditLogListResponse,
)
def admin_list_audit_logs(
    username: str | None = None,
    action: str | None = None,
    entity_type: str | None = None,
    search: str | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    return crud.list_audit_logs(
        db=db,
        username=username,
        action=action,
        entity_type=entity_type,
        search=search,
        page=page,
        page_size=page_size,
    )

@app.get("/api/admin/students", response_model=schemas.StudentListResponse)
def admin_list_students(
    search: str | None = None,
    graduation_status: str | None = None,
    is_deleted: bool = Query(default=False),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    return crud.list_students(
        db=db,
        search=search,
        graduation_status=graduation_status,
        is_deleted=is_deleted,
        page=page,
        page_size=page_size,
    )


@app.post("/api/admin/students/validate")
def validate_admin_student_data(
    payload: dict = Body(...),
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    errors = []
    current_student_id = str(
        payload.get("current_student_id")
        or payload.get("original_student_id")
        or ""
    ).strip() or None

    student_id = str(payload.get("student_id") or "").strip()

    if student_id:
        existing_student = crud.get_student(db, student_id)

        if existing_student and not is_same_student_id(
            existing_student.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "student_id"],
                    "msg": "Student already exists",
                    "type": "value_error.duplicate",
                }
            )

    email = str(payload.get("email") or "").strip()

    if email:
        existing_email = crud.get_student_by_email(db, email)

        if existing_email and not is_same_student_id(
            existing_email.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "email"],
                    "msg": "Email already exists",
                    "type": "value_error.duplicate",
                }
            )

    citizen_id = str(payload.get("citizen_id") or "").strip()

    if citizen_id and len(citizen_id) == 12 and citizen_id.isdigit():
        citizen_id_hash = crud.calculate_citizen_id_hash(citizen_id)
        existing_citizen = crud.get_student_by_citizen_id_hash(
            db,
            citizen_id_hash,
        )

        if existing_citizen and not is_same_student_id(
            existing_citizen.student_id,
            current_student_id,
        ):
            errors.append(
                {
                    "loc": ["body", "citizen_id"],
                    "msg": "Citizen ID already exists",
                    "type": "value_error.duplicate",
                }
            )

    return {
        "valid": len(errors) == 0,
        "errors": errors,
    }


@app.post(
    "/api/admin/students",
    response_model=schemas.StudentResponse,
)
def admin_create_student(
    student_data: schemas.StudentCreate,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    normalized_faculty_name = normalize_faculty_name(
        student_data.faculty_name
    )

    is_graduated = student_data.graduation_status == "GRADUATED"

    update_values = {
        "faculty_name": normalized_faculty_name,
    }

    if is_graduated:
        missing_fields = []

        if student_data.gpa is None:
            missing_fields.append("GPA")

        if student_data.graduation_date is None:
            missing_fields.append("Ngày tốt nghiệp")

        if student_data.graduation_year is None:
            missing_fields.append("Năm tốt nghiệp")

        if missing_fields:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Sinh viên đã tốt nghiệp còn thiếu thông tin: "
                    f"{', '.join(missing_fields)}."
                ),
            )

        update_values["classification"] = classification_from_gpa(
            student_data.gpa
        )
    else:
        # Sinh viên chưa tốt nghiệp chỉ được lưu trong MySQL.
        update_values.update(
            {
                "gpa": None,
                "classification": None,
                "graduation_date": None,
                "graduation_year": None,
            }
        )

    student_data = student_data.model_copy(update=update_values)

    validate_unique_student_fields(db, student_data)

    student = crud.create_student(
        db=db,
        student_data=student_data,
        updated_by=current_admin.username,
        commit=False,
    )

    # A direct graduate record is saved in MySQL first.  It reaches Fabric only
    # after this administrator explicitly approves it from the Admin page.
    if is_graduated:
        student.blockchain_tx_id = None
        student.is_mismatch = False
        student.approval_status = "PENDING_BLOCKCHAIN"
        student.approved_by = None
        student.approved_at = None
        audit_details = (
            f"Đã tạo sinh viên {student.student_id}. "
            "Hồ sơ tốt nghiệp đang chờ duyệt ghi Blockchain."
        )
    else:
        student.metadata_hash = None
        student.blockchain_tx_id = None
        student.is_mismatch = False
        student.approval_status = "APPROVED"
        student.approved_by = current_admin.username
        student.approved_at = now_utc()
        audit_details = (
            f"Đã tạo sinh viên {student.student_id}. "
            "Trạng thái: Chưa tốt nghiệp."
        )

    student.updated_by = current_admin.username

    crud.create_audit_log(
        db=db,
        username=current_admin.username,
        action="CREATE_STUDENT",
        entity_type="Student",
        entity_id=student.student_id,
        details=audit_details,
        commit=False,
    )

    try:
        db.commit()
        db.refresh(student)
    except Exception:
        db.rollback()
        logger.exception(
            "MySQL commit failed when creating student %s",
            student.student_id,
        )
        raise HTTPException(
            status_code=500,
            detail="Không thể lưu thông tin sinh viên vào cơ sở dữ liệu.",
        )

    logger.info(
        "Student %s created by %s. Pending blockchain approval: %s",
        student.student_id,
        current_admin.username,
        is_graduated,
    )

    return student



STUDENT_AUDIT_FIELD_LABELS = {
    "full_name": "Họ và tên",
    "date_of_birth": "Ngày sinh",
    "email": "Email",
    "institution_name": "Trường",
    "faculty_name": "Khoa",
    "major": "Chuyên ngành",
    "training_mode": "Hình thức đào tạo",
    "degree_id": "Mã bằng",
    "degree_type": "Loại bằng",
    "entrance_year": "Khóa nhập học",
    "graduation_status": "Trạng thái tốt nghiệp",
    "graduation_date": "Ngày tốt nghiệp",
    "graduation_year": "Năm tốt nghiệp",
    "classification": "Xếp loại",
    "gpa": "GPA",
    "expected_graduation_date": "Ngày tốt nghiệp dự kiến",
    "citizen_id": "Số CCCD/CMND",
}


STUDENT_AUDIT_VALUE_LABELS = {
    "GRADUATED": "Đã tốt nghiệp",
    "NOT_GRADUATED": "Chưa tốt nghiệp",
    "Excellent": "Xuất sắc",
    "Very Good": "Giỏi",
    "Good": "Khá",
    "Average": "Trung bình",
    "Weak": "Yếu",
    "Bachelor": "Cử nhân",
}


def format_student_audit_value(field_name: str, value) -> str:
    if value is None or value == "":
        return "Để trống"

    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M:%S")

    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")

    if field_name == "gpa":
        numeric_value = float(value)
        if numeric_value.is_integer():
            return f"{numeric_value:.1f}"
        return f"{numeric_value:.2f}".rstrip("0")

    return STUDENT_AUDIT_VALUE_LABELS.get(str(value), str(value))


def student_audit_values_equal(field_name: str, old_value, new_value) -> bool:
    if field_name == "gpa":
        if old_value is None or new_value is None:
            return old_value is new_value

        try:
            return float(old_value) == float(new_value)
        except (TypeError, ValueError):
            return str(old_value) == str(new_value)

    return old_value == new_value


def build_student_change_details(
    old_values: dict,
    new_values: dict,
) -> tuple[str, set[str]]:
    changes = []
    changed_fields = set()

    # Dat xep loai truoc GPA de noi dung lich su de doc hon.
    field_names = list(new_values.keys())
    if "classification" in field_names and "gpa" in field_names:
        field_names.remove("classification")
        field_names.insert(field_names.index("gpa"), "classification")

    for field_name in field_names:
        new_value = new_values[field_name]
        label = STUDENT_AUDIT_FIELD_LABELS.get(field_name)

        if not label:
            continue

        if field_name == "citizen_id":
            if new_value:
                changes.append("Số CCCD/CMND: đã thay đổi")
                changed_fields.add(field_name)
            continue

        old_value = old_values.get(field_name)

        if student_audit_values_equal(field_name, old_value, new_value):
            continue

        old_display = format_student_audit_value(field_name, old_value)
        new_display = format_student_audit_value(field_name, new_value)
        changes.append(
            f'{label}: từ "{old_display}" thành "{new_display}"'
        )
        changed_fields.add(field_name)

    return "; ".join(changes), changed_fields

@app.put(
    "/api/admin/students/{student_id}",
    response_model=schemas.StudentResponse,
)
def admin_update_student(
    student_id: str,
    student_data: schemas.StudentUpdate,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    # ADMIN_STUDENT_IMMUTABLE_API_V1
    raise HTTPException(
        status_code=405,
        detail="Student records are immutable; update and delete are disabled.",
    )

    student = crud.get_student(db, student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    update_data = student_data.model_dump(exclude_unset=True)
    duplicate_errors = []

    if "faculty_name" in update_data:
        normalized_faculty_name = normalize_faculty_name(
            update_data["faculty_name"]
        )
        student_data = student_data.model_copy(
            update={"faculty_name": normalized_faculty_name}
        )
        update_data["faculty_name"] = normalized_faculty_name

    if (
        student.graduation_status == "GRADUATED"
        and update_data.get("graduation_status") == "NOT_GRADUATED"
    ):
        raise HTTPException(
            status_code=400,
            detail="Graduated student cannot be changed to not graduated",
        )

    if "gpa" in update_data:
        calculated_classification = classification_from_gpa(update_data["gpa"])

        if calculated_classification:
            student_data = student_data.model_copy(
                update={"classification": calculated_classification}
            )
            update_data["classification"] = calculated_classification

    target_graduation_status = update_data.get(
        "graduation_status",
        student.graduation_status,
    )

    if target_graduation_status == "GRADUATED":
        graduation_values = {
            "GPA": update_data.get("gpa", student.gpa),
            "Ngày tốt nghiệp": update_data.get(
                "graduation_date",
                student.graduation_date,
            ),
            "Năm tốt nghiệp": update_data.get(
                "graduation_year",
                student.graduation_year,
            ),
        }
        missing_fields = [
            label
            for label, value in graduation_values.items()
            if value is None
        ]

        if missing_fields:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Chưa thể xác nhận tốt nghiệp vì còn thiếu: "
                    f"{', '.join(missing_fields)}."
                ),
            )
    else:
        # Sinh viên chưa tốt nghiệp chưa có dữ liệu tốt nghiệp chính thức.
        empty_graduation_values = {
            "gpa": None,
            "classification": None,
            "graduation_date": None,
            "graduation_year": None,
        }
        student_data = student_data.model_copy(
            update=empty_graduation_values
        )
        update_data.update(empty_graduation_values)

    if "degree_id" in update_data and update_data["degree_id"]:
        existing_degree = crud.get_student_by_degree_id(
            db,
            update_data["degree_id"],
        )

        if existing_degree and not is_same_student_id(
            existing_degree.student_id,
            student_id,
        ):
            duplicate_errors.append(
                {
                    "loc": ["body", "degree_id"],
                    "msg": "Degree already exists",
                    "type": "value_error.duplicate",
                }
            )

    if "email" in update_data and update_data["email"]:
        existing_email = crud.get_student_by_email(db, update_data["email"])

        if existing_email and not is_same_student_id(
            existing_email.student_id,
            student_id,
        ):
            duplicate_errors.append(
                {
                    "loc": ["body", "email"],
                    "msg": "Email already exists",
                    "type": "value_error.duplicate",
                }
            )

    if "citizen_id" in update_data and update_data["citizen_id"]:
        citizen_id_hash = crud.calculate_citizen_id_hash(
            update_data["citizen_id"]
        )

        existing_citizen = crud.get_student_by_citizen_id_hash(
            db,
            citizen_id_hash,
        )

        if existing_citizen and not is_same_student_id(
            existing_citizen.student_id,
            student_id,
        ):
            duplicate_errors.append(
                {
                    "loc": ["body", "citizen_id"],
                    "msg": "Citizen ID already exists",
                    "type": "value_error.duplicate",
                }
            )

    raise_duplicate_field_errors(duplicate_errors)

    old_values = {
        field_name: getattr(student, field_name, None)
        for field_name in update_data
        if field_name != "citizen_id"
    }
    change_details, changed_fields = build_student_change_details(
        old_values,
        update_data,
    )

    had_verification_record = bool(student.blockchain_tx_id)
    requires_blockchain_approval = (
        target_graduation_status == "GRADUATED"
        and not had_verification_record
    )
    updated_student = crud.update_student(
        db=db,
        student=student,
        student_data=student_data,
        updated_by=current_admin.username,
        sync_hash=target_graduation_status == "GRADUATED",
        commit=False,
    )

    verification_status = (
        "Pending Approval" if requires_blockchain_approval else "Not Synced"
    )
    sync_record = None

    if target_graduation_status == "GRADUATED" and not requires_blockchain_approval:
        try:
            verification_result = verify_student_on_blockchain(updated_student)
            verification_status = verification_result["verification_status"]
            updated_student.is_mismatch = verification_status == "Mismatch"
        except Exception:
            db.rollback()
            logger.exception(
                "Student verification failed when updating %s",
                student_id,
            )
            raise HTTPException(
                status_code=503,
                detail=(
                    "Hệ thống xác minh dữ liệu đang tạm thời gián đoạn. "
                    "Thông tin sinh viên chưa được thay đổi."
                ),
            )
    elif target_graduation_status == "GRADUATED":
        # A newly-graduated or legacy unsynced record must await the personal
        # admin decision instead of silently generating a Fabric transaction.
        updated_student.blockchain_tx_id = None
        updated_student.is_mismatch = False
        updated_student.approval_status = "PENDING_BLOCKCHAIN"
        updated_student.approved_by = None
        updated_student.approved_at = None
    else:
        updated_student.metadata_hash = None
        updated_student.blockchain_tx_id = None
        updated_student.is_mismatch = False
        updated_student.approval_status = "APPROVED"
        updated_student.approved_by = None
        updated_student.approved_at = None

    # Khong tao lich su neu admin bam cap nhat nhung du lieu khong doi.
    if change_details:
        audit_details = f"Đã cập nhật: {change_details}."

        # Chi hien ket qua doi chieu khi GPA thuc su thay doi.
        if "gpa" in changed_fields:
            verification_status_vi = {
                "Verified": "Khớp",
                "Mismatch": "Không khớp",
                "Not Found": "Chưa có dữ liệu xác minh",
                "Not Synced": "Chưa cần xác minh",
                "Pending Approval": "Chờ duyệt ghi Blockchain",
            }.get(verification_status, verification_status)
            audit_details += (
                " Kết quả đối chiếu dữ liệu: "
                f"{verification_status_vi}."
            )

        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="UPDATE_STUDENT",
            entity_type="Student",
            entity_id=student.student_id,
            details=audit_details,
            commit=False,
        )

    try:
        db.commit()
        db.refresh(updated_student)
    except Exception:
        db.rollback()

        if sync_record is not None:
            logger.critical(
                "Fabric synced updated student %s but database commit "
                "failed. Tx: %s",
                student_id,
                sync_record.tx_id,
            )

        logger.exception(
            "Failed to save updated student %s",
            student_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Database could not save the updated student record.",
        )

    logger.info(
        "Student %s updated by %s; blockchain verification status: %s",
        student_id,
        current_admin.username,
        verification_status,
    )

    return updated_student





@app.post(
    "/api/admin/students/{student_id}/reject-blockchain",
    response_model=schemas.ApprovalResponse,
)
def admin_reject_student_blockchain(
    student_id: str,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    # MULTI_USER_ROW_LOCK_V1: serialize decisions for this student only.
    student = crud.get_student_for_update(db, student_id)

    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    if student.approval_status != "PENDING_BLOCKCHAIN":
        raise HTTPException(
            status_code=409,
            detail="Student is not waiting for blockchain approval.",
        )

    try:
        decision = record_student_rejection_on_blockchain(student)
    except Exception:
        db.rollback()
        logger.exception(
            "Fabric rejection transaction failed for student %s",
            student_id,
        )
        raise HTTPException(
            status_code=503,
            detail=(
                "Blockchain is temporarily unavailable. "
                "Student remains pending."
            ),
        )

    student.approval_status = "REJECTED_BLOCKCHAIN"
    student.approved_by = current_admin.username
    student.approved_at = now_utc()
    student.updated_by = current_admin.username
    student.blockchain_tx_id = decision.tx_id
    student.is_mismatch = False

    crud.create_audit_log(
        db=db,
        username=current_admin.username,
        action="REJECT_STUDENT_BLOCKCHAIN",
        entity_type="Student",
        entity_id=student.student_id,
        details=(
            "Đã từ chối duyệt và ghi quyết định lên Blockchain. "
            f"Mã giao dịch: {decision.tx_id or 'đã tạo'}"
        ),
        commit=False,
    )

    try:
        db.commit()
        db.refresh(student)
    except Exception:
        db.rollback()
        logger.exception(
            "Fabric rejection was recorded but MySQL update failed for student %s",
            student_id,
        )
        raise HTTPException(
            status_code=500,
            detail="Blockchain decision was recorded but could not be saved locally.",
        )

    return schemas.ApprovalResponse(
        message="Student rejection decision was recorded on blockchain.",
        student_id=student.student_id,
        approval_status=student.approval_status,
        approved_by=student.approved_by,
        approved_at=student.approved_at,
        blockchain_tx_id=student.blockchain_tx_id,
        metadata_hash=student.metadata_hash,
    )

@app.post(
    "/api/admin/students/{student_id}/approve-blockchain",
    response_model=schemas.ApprovalResponse,
)
def admin_approve_student_blockchain(
    student_id: str,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    # MULTI_USER_ROW_LOCK_V1: serialize decisions for this student only.
    student = crud.get_student_for_update(db, student_id)

    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    if student.approval_status != "PENDING_BLOCKCHAIN":
        raise HTTPException(
            status_code=409,
            detail="Student is not waiting for blockchain approval.",
        )

    if student.graduation_status != "GRADUATED":
        raise HTTPException(
            status_code=409,
            detail="Only graduated students can be recorded on blockchain.",
        )

    try:
        record = sync_student_to_blockchain(student)
    except Exception:
        db.rollback()
        logger.exception(
            "Blockchain sync failed while approving student %s",
            student_id,
        )
        raise HTTPException(
            status_code=503,
            detail="Blockchain is temporarily unavailable. Student remains pending.",
        )

    student.metadata_hash = record.metadata_hash
    student.metadata_hash_version = record.hash_version
    student.blockchain_tx_id = record.tx_id
    student.is_mismatch = False
    student.approval_status = "APPROVED"
    student.approved_by = current_admin.username
    student.approved_at = now_utc()
    student.updated_by = current_admin.username

    crud.create_audit_log(
        db=db,
        username=current_admin.username,
        action="APPROVE_STUDENT_BLOCKCHAIN",
        entity_type="Student",
        entity_id=student.student_id,
        details=(
            "Đã duyệt ghi hồ sơ tốt nghiệp lên Blockchain. "
            f"Mã giao dịch Fabric: {record.tx_id}."
        ),
        commit=False,
    )

    try:
        db.commit()
        db.refresh(student)
    except Exception:
        db.rollback()
        logger.critical(
            "Fabric synced approved student %s but MySQL commit failed. Tx: %s",
            student_id,
            record.tx_id,
        )
        raise HTTPException(
            status_code=500,
            detail="Fabric recorded the transaction but MySQL could not save its receipt.",
        )

    return schemas.ApprovalResponse(
        message="Student approved and recorded on blockchain.",
        student_id=student.student_id,
        approval_status=student.approval_status,
        approved_by=student.approved_by,
        approved_at=student.approved_at,
        blockchain_tx_id=student.blockchain_tx_id,
        metadata_hash=student.metadata_hash,
    )

# ADMIN_STUDENT_PRIVATE_DETAIL_V1
@app.get("/api/admin/students/{student_id}", response_model=schemas.AdminStudentDetailResponse)
def admin_get_student(
    student_id: str,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    student = crud.get_student(db, student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    try:
        citizen_id = decrypt_value(student.citizen_id_encrypted)
    except EncryptionError as exc:
        logger.exception("Cannot decrypt citizen ID for student %s", student_id)
        raise HTTPException(
            status_code=500,
            detail="Cannot decrypt student citizen ID",
        ) from exc

    response = schemas.AdminStudentDetailResponse.model_validate(student)
    return response.model_copy(update={"citizen_id": citizen_id})



@app.delete("/api/admin/students/{student_id}")
def admin_delete_student(
    student_id: str,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    # ADMIN_STUDENT_IMMUTABLE_API_V1
    raise HTTPException(
        status_code=405,
        detail="Student records are immutable; update and delete are disabled.",
    )

    student = crud.get_student(db, student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    try:
        crud.delete_student(
            db=db,
            student=student,
            deleted_by=current_admin.username,
            commit=False,
        )

        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="DELETE_STUDENT",
            entity_type="Student",
            entity_id=student.student_id,
            details=f"Đã xóa mềm sinh viên {student.student_id}.",
            commit=False,
        )

        db.commit()
    except Exception:
        db.rollback()

        logger.exception(
            "Failed to soft-delete student %s",
            student_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Unable to delete student",
        )

    logger.info(
        "Student %s soft-deleted by %s",
        student_id,
        current_admin.username,
    )

    return {
        "message": "Student deleted",
        "student_id": student_id,
    }

@app.post(
    "/api/admin/students/{student_id}/restore",
    response_model=schemas.StudentResponse,
)
def admin_restore_student(
    student_id: str,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    student = crud.get_deleted_student(db, student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Deleted student not found",
        )

    try:
        restored_student = crud.restore_student(
            db=db,
            student=student,
            commit=False,
        )

        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="RESTORE_STUDENT",
            entity_type="Student",
            entity_id=restored_student.student_id,
            details=(
                f"Đã khôi phục sinh viên {restored_student.student_id} "
                "từ danh sách đã xóa."
            ),
            commit=False,
        )

        db.commit()
        db.refresh(restored_student)
    except Exception:
        db.rollback()

        logger.exception(
            "Failed to restore student %s",
            student_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Unable to restore student",
        )

    logger.info(
        "Student %s restored by %s",
        student_id,
        current_admin.username,
    )

    return restored_student




@app.post(
    "/api/external-requests",
    response_model=schemas.ExternalRequestResponse,
)
def create_external_request(
    request_data: schemas.ExternalRequestCreate,
    db: Session = Depends(get_db),
):
    external_request = crud.create_external_request(
        db=db,
        request_data=request_data,
    )

    logger.info(
        "External request %s created for student %s",
        external_request.id,
        external_request.student_id,
    )
    return external_request


@app.post(
    "/api/admin/external-requests/import-file",
    response_model=schemas.ExternalFileImportResponse,
    status_code=status.HTTP_201_CREATED,
)
async def import_external_student_file(
    company_name: str = Form(..., min_length=2, max_length=100),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    """Receive student records from an external source for admin review."""
    original_filename = Path(file.filename or "external-student-data").name
    extension = Path(original_filename).suffix.lower()

    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(
            status_code=400,
            detail="Chỉ chấp nhận tệp CSV hoặc XLSX.",
        )

    file_content = await file.read()
    if not file_content:
        raise HTTPException(
            status_code=400,
            detail="Tệp tải lên đang trống.",
        )
    if len(file_content) > MAX_EXTERNAL_IMPORT_BYTES:
        raise HTTPException(
            status_code=413,
            detail="Tệp tải lên không được vượt quá 5 MB.",
        )

    rows = parse_external_student_file(
        original_filename,
        file_content,
    )
    if not rows:
        raise HTTPException(
            status_code=400,
            detail="Tệp không có dòng dữ liệu sinh viên.",
        )

    imported_rows = 0
    ready_for_review_count = 0
    invalid_count = 0
    duplicate_count = 0
    non_fpt_count = 0
    invalid_rows = []
    seen_student_ids = set()
    seen_emails = set()
    seen_citizen_ids = set()

    try:
        for row in rows:
            imported_data = row["data"]
            student_id = str(
                imported_data.get("student_id") or ""
            ).strip().upper()

            if not student_id:
                invalid_count += 1
                invalid_rows.append(
                    {
                        "row_number": row["row_number"],
                        "student_id": None,
                        "message": "Dòng dữ liệu thiếu mã số sinh viên.",
                    }
                )
                continue

            student_data, verification_status, verification_message = (
                prepare_external_student(db, imported_data)
            )

            email_key = str(imported_data.get("email") or "").strip().lower()
            citizen_key = str(
                imported_data.get("citizen_id") or ""
            ).strip()
            duplicate_labels = []
            if student_id in seen_student_ids:
                duplicate_labels.append("mã số sinh viên")
            if email_key and email_key in seen_emails:
                duplicate_labels.append("email")
            if citizen_key and citizen_key in seen_citizen_ids:
                duplicate_labels.append("số CCCD/CMND")

            if duplicate_labels:
                student_data = None
                verification_status = "DUPLICATE_IN_FILE"
                verification_message = (
                    "Dòng dữ liệu bị trùng "
                    + ", ".join(duplicate_labels)
                    + " trong cùng tệp."
                )

            seen_student_ids.add(student_id)
            if email_key:
                seen_emails.add(email_key)
            if citizen_key:
                seen_citizen_ids.add(citizen_key)

            if verification_status == "READY_FOR_REVIEW":
                ready_for_review_count += 1
            elif verification_status == "NOT_FPT_INSTITUTION":
                non_fpt_count += 1
            elif verification_status.startswith("DUPLICATE"):
                duplicate_count += 1
            else:
                invalid_count += 1

            # Gửi đầy đủ lý do của mọi dòng không thể duyệt về giao diện.
            # Trước đây danh sách này chỉ có trường hợp thiếu MSSV nên admin
            # không biết các dòng GPA sai, trùng dữ liệu hoặc không thuộc FPT.
            if verification_status != "READY_FOR_REVIEW":
                invalid_rows.append(
                    {
                        "row_number": row["row_number"],
                        "student_id": student_id or None,
                        "message": verification_message,
                    }
                )

            crud.create_external_request(
                db=db,
                request_data={
                    "student_id": student_id,
                    "full_name": imported_data.get("full_name"),
                    "gpa": parse_external_decimal(imported_data.get("gpa")),
                    "graduation_status": normalize_external_status(
                        imported_data.get("graduation_status")
                    ),
                    "requester_name": company_name.strip(),
                    "requester_email": None,
                    "requester_type": "EXTERNAL_STUDENT_FILE",
                    "message": None,
                    "source_file_name": original_filename,
                    "source_row_number": row["row_number"],
                    "verification_status": verification_status,
                    "verification_message": verification_message,
                    "matched_student_name": None,
                    "import_data": json.dumps(
                        imported_data,
                        ensure_ascii=False,
                        default=str,
                    ),
                },
                commit=False,
            )
            imported_rows += 1

        if imported_rows == 0:
            db.rollback()
            raise HTTPException(
                status_code=422,
                detail="Không có hồ sơ sinh viên nào có mã số để tiếp nhận.",
            )

        rows_needing_review = (
            invalid_count + duplicate_count + non_fpt_count
        )
        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="IMPORT_EXTERNAL_STUDENT_FILE",
            entity_type="ExternalRequest",
            entity_id=original_filename,
            details=(
                f"Đã ghi nhận tệp {original_filename} từ nguồn "
                f"{company_name.strip()}: {imported_rows} dòng đã ghi nhận; "
                f"{ready_for_review_count} hồ sơ có thể duyệt; "
                f"{rows_needing_review} dòng cần kiểm tra."
            ),
            commit=False,
        )
        db.commit()
    except HTTPException:
        raise
    except Exception:
        db.rollback()
        logger.exception(
            "Failed to import external student file %s",
            original_filename,
        )
        raise HTTPException(
            status_code=500,
            detail="Không thể tiếp nhận tệp dữ liệu sinh viên.",
        )

    logger.info(
        "External student file %s imported by %s",
        original_filename,
        current_admin.username,
    )
    return {
        "file_name": original_filename,
        "total_rows": len(rows),
        "imported_rows": imported_rows,
        "ready_for_review_count": ready_for_review_count,
        "invalid_count": invalid_count,
        "duplicate_count": duplicate_count,
        "non_fpt_count": non_fpt_count,
        "invalid_rows": invalid_rows,
    }


@app.get(
    "/api/admin/external-requests",
    response_model=schemas.ExternalRequestListResponse,
)
def admin_list_external_requests(
    status: str | None = None,
    search: str | None = None,
    is_deleted: bool = Query(default=False),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    return crud.list_external_requests(
        db=db,
        status=status,
        search=search,
        page=page,
        page_size=page_size,
        is_deleted=is_deleted,
    )


@app.delete(
    "/api/admin/external-requests/{request_id}",
    response_model=schemas.ExternalRequestResponse,
)
def admin_archive_external_request(
    request_id: int,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    external_request = crud.get_external_request(db, request_id)

    if not external_request:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy hồ sơ nguồn ngoài.",
        )

    if external_request.status not in {"APPROVED", "REJECTED"}:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "Chỉ có thể lưu trữ hồ sơ đã duyệt hoặc đã từ chối. "
                "Hồ sơ đang chờ xử lý phải được duyệt hoặc từ chối trước."
            ),
        )

    try:
        archived_request = crud.archive_external_request(
            db=db,
            external_request=external_request,
            deleted_by=current_admin.username,
            commit=False,
        )
        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="ARCHIVE_EXTERNAL_REQUEST",
            entity_type="ExternalRequest",
            entity_id=archived_request.id,
            details=(
                "Đã lưu trữ hồ sơ nguồn ngoài của sinh viên "
                f"{archived_request.student_id}. Dữ liệu sinh viên và "
                "Blockchain không bị thay đổi."
            ),
            commit=False,
        )
        db.commit()
        db.refresh(archived_request)
    except Exception:
        db.rollback()
        logger.exception(
            "Failed to archive external request %s",
            request_id,
        )
        raise HTTPException(
            status_code=500,
            detail="Không thể lưu trữ hồ sơ nguồn ngoài.",
        )

    return archived_request


@app.post(
    "/api/admin/external-requests/{request_id}/restore",
    response_model=schemas.ExternalRequestResponse,
)
def admin_restore_external_request(
    request_id: int,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    external_request = crud.get_deleted_external_request(db, request_id)

    if not external_request:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy hồ sơ đã lưu trữ.",
        )

    try:
        restored_request = crud.restore_external_request(
            db=db,
            external_request=external_request,
            commit=False,
        )
        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="RESTORE_EXTERNAL_REQUEST",
            entity_type="ExternalRequest",
            entity_id=restored_request.id,
            details=(
                "Đã khôi phục hồ sơ nguồn ngoài của sinh viên "
                f"{restored_request.student_id}."
            ),
            commit=False,
        )
        db.commit()
        db.refresh(restored_request)
    except Exception:
        db.rollback()
        logger.exception(
            "Failed to restore external request %s",
            request_id,
        )
        raise HTTPException(
            status_code=500,
            detail="Không thể khôi phục hồ sơ nguồn ngoài.",
        )

    return restored_request


@app.post(
    "/api/admin/external-requests/{request_id}/reject",
    response_model=schemas.ExternalRequestResponse,
)
def admin_reject_external_request(
    request_id: int,
    reject_data: schemas.ExternalRequestReject,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    # MULTI_USER_ROW_LOCK_V1: prevent simultaneous accept/reject decisions.
    external_request = crud.get_external_request_for_update(db, request_id)
    if not external_request:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy hồ sơ nguồn ngoài.",
        )
    if external_request.status != "PENDING_REVIEW":
        raise HTTPException(
            status_code=400,
            detail="Hồ sơ này đã được xử lý.",
        )

    # EXTERNAL_REQUEST_FABRIC_DECISION_V1
    source_data_hash = calculate_external_request_source_hash(external_request)
    try:
        decision = record_external_request_rejection_on_blockchain(
            external_request.student_id,
            source_data_hash,
        )
    except Exception as exc:
        db.rollback()
        logger.exception(
            "Failed to record external rejection on Fabric for request %s",
            request_id,
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "BLOCKCHAIN_REJECTION_UNAVAILABLE",
                "message": (
                    "Hệ thống chưa thể ghi quyết định từ chối lên "
                    "Blockchain. Hồ sơ chưa bị thay đổi; vui lòng thử lại."
                ),
            },
        ) from exc

    try:
        updated_request = crud.reject_external_request(
            db=db,
            external_request=external_request,
            reject_reason=reject_data.reject_reason,
            reviewed_by=current_admin.username,
            commit=False,
        )
        updated_request.source_data_hash = source_data_hash
        updated_request.blockchain_tx_id = decision.get("transaction_id")
        updated_request.blockchain_action = "EXTERNAL_REQUEST_REJECTED"
        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="REJECT_EXTERNAL_STUDENT_IMPORT",
            entity_type="ExternalRequest",
            entity_id=updated_request.id,
            details=(
                "Đã từ chối hồ sơ nguồn ngoài của sinh viên "
                f"{updated_request.student_id}. Mã giao dịch Fabric: "
                f"{updated_request.blockchain_tx_id or 'không đọc được'}."
            ),
            commit=False,
        )
        db.commit()
        db.refresh(updated_request)
    except Exception:
        db.rollback()
        logger.critical(
            "Fabric rejected external request %s but database commit failed. "
            "Tx: %s",
            request_id,
            decision.get("transaction_id"),
            exc_info=True,
        )
        raise HTTPException(
            status_code=500,
            detail="Blockchain đã ghi quyết định nhưng MySQL chưa lưu được kết quả.",
        )

    return updated_request


@app.post(
    "/api/admin/external-requests/{request_id}/approve",
    response_model=schemas.ExternalRequestResponse,
)
def admin_approve_external_request(
    request_id: int,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    # MULTI_USER_ROW_LOCK_V1: prevent simultaneous accept/reject decisions.
    external_request = crud.get_external_request_for_update(db, request_id)
    if not external_request:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy hồ sơ nguồn ngoài.",
        )
    if external_request.status != "PENDING_REVIEW":
        raise HTTPException(
            status_code=400,
            detail="Hồ sơ này đã được xử lý.",
        )

    try:
        imported_data = json.loads(external_request.import_data or "{}")
    except (TypeError, json.JSONDecodeError):
        imported_data = {}
    if not isinstance(imported_data, dict):
        imported_data = {}

    imported_data.setdefault("student_id", external_request.student_id)
    imported_data.setdefault("full_name", external_request.full_name)
    imported_data.setdefault(
        "gpa",
        str(external_request.gpa)
        if external_request.gpa is not None
        else None,
    )
    imported_data.setdefault(
        "graduation_status",
        external_request.graduation_status,
    )

    student_data, verification_status, verification_message = (
        prepare_external_student(db, imported_data)
    )
    external_request.verification_status = verification_status
    external_request.verification_message = verification_message

    if student_data is None or verification_status != "READY_FOR_REVIEW":
        db.commit()
        raise HTTPException(
            status_code=409,
            detail={
                "code": verification_status,
                "message": verification_message,
            },
        )

    is_graduated = student_data.graduation_status == "GRADUATED"
    student = crud.create_student(
        db=db,
        student_data=student_data,
        updated_by=current_admin.username,
        commit=False,
    )

    # EXTERNAL_STUDENT_INTAKE_PENDING_BLOCKCHAIN_V1
    # Tiếp nhận hồ sơ nguồn ngoài chỉ tạo bản ghi MySQL. Hồ sơ tốt nghiệp
    # phải được duyệt lần cuối trong Quản lý sinh viên trước khi gọi Fabric.
    student.blockchain_tx_id = None
    student.is_mismatch = False

    if is_graduated:
        student.approval_status = "PENDING_BLOCKCHAIN"
        student.approved_by = None
        student.approved_at = None
    else:
        # Sinh viên chưa tốt nghiệp chưa thuộc phạm vi ghi nhận tốt nghiệp.
        student.metadata_hash = None
        student.approval_status = "APPROVED"
        student.approved_by = current_admin.username
        student.approved_at = now_utc()

    student.updated_by = current_admin.username

    updated_request = crud.approve_external_request(
        db=db,
        external_request=external_request,
        reviewed_by=current_admin.username,
        commit=False,
    )
    updated_request.verification_status = "IMPORTED_TO_SCHOOL"
    if is_graduated:
        updated_request.verification_message = (
            "Đã tiếp nhận sinh viên vào Quản lý sinh viên. "
            "Hồ sơ tốt nghiệp đang chờ duyệt ghi Blockchain."
        )
        audit_details = (
            "Đã tiếp nhận hồ sơ nguồn ngoài của sinh viên "
            f"{student.student_id} vào Quản lý sinh viên. "
            "Chưa tạo giao dịch Fabric; hồ sơ đang chờ duyệt."
        )
    else:
        updated_request.verification_message = (
            "Đã tiếp nhận sinh viên chưa tốt nghiệp vào Quản lý sinh viên. "
            "Hồ sơ chưa thuộc phạm vi ghi nhận tốt nghiệp trên Blockchain."
        )
        audit_details = (
            "Đã tiếp nhận hồ sơ nguồn ngoài của sinh viên chưa tốt nghiệp "
            f"{student.student_id} vào Quản lý sinh viên."
        )

    updated_request.matched_student_name = student.full_name
    updated_request.blockchain_tx_id = None
    updated_request.blockchain_action = (
        "STUDENT_PENDING_BLOCKCHAIN"
        if is_graduated
        else "STUDENT_ACCEPTED_NOT_GRADUATED"
    )

    crud.create_audit_log(
        db=db,
        username=current_admin.username,
        action="ACCEPT_EXTERNAL_STUDENT_IMPORT",
        entity_type="ExternalRequest",
        entity_id=updated_request.id,
        details=audit_details,
        commit=False,
    )

    try:
        db.commit()
        db.refresh(student)
        db.refresh(updated_request)
    except Exception:
        db.rollback()
        logger.exception(
            "Database commit failed while accepting external student %s",
            student.student_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Không thể lưu hồ sơ sinh viên đã tiếp nhận.",
        )

    logger.info(
        "External student request %s accepted; student %s created",
        request_id,
        student.student_id,
    )
    return updated_request

@app.post(
    "/api/certificate-requests/otp/send",
    response_model=schemas.CertificateOtpSendResponse,
)
def send_certificate_request_otp(
    otp_data: schemas.CertificateOtpSendRequest,
    db: Session = Depends(get_db),
):
    student = crud.get_student(db, otp_data.student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    if student.graduation_status != "GRADUATED":
        raise HTTPException(
            status_code=400,
            detail="Student has not graduated yet",
        )

    if crud.has_recent_certificate_otp(
        db=db,
        student_id=otp_data.student_id,
        requester_email=otp_data.requester_email,
    ):
        raise HTTPException(
            status_code=status.HTTP_429_TOO_MANY_REQUESTS,
            detail="Please wait 60 seconds before requesting another OTP",
        )

    otp_code = crud.generate_certificate_otp()
    expires_at = crud.get_certificate_otp_expiration()

    verification = crud.create_certificate_email_verification(
        db=db,
        student_id=otp_data.student_id,
        requester_email=otp_data.requester_email,
        otp_code=otp_code,
        expires_at=expires_at,
    )

    try:
        send_certificate_request_otp_email(
            db=db,
            to_email=otp_data.requester_email,
            student_id=otp_data.student_id,
            otp_code=otp_code,
        )
    except Exception:
        try:
            db.delete(verification)
            db.commit()
        except Exception:
            db.rollback()

        logger.exception(
            "Failed to send certificate OTP for student %s",
            otp_data.student_id,
        )

        raise HTTPException(
            status_code=503,
            detail="Unable to send OTP email. Please try again later.",
        )

    return schemas.CertificateOtpSendResponse(
        message="OTP has been sent to your email",
        expires_in_seconds=crud.OTP_EXPIRY_MINUTES * 60,
    )

@app.post(
    "/api/certificate-requests/otp/verify",
    response_model=schemas.CertificateOtpVerifyResponse,
)
def verify_certificate_request_otp(
    otp_data: schemas.CertificateOtpVerifyRequest,
    db: Session = Depends(get_db),
):
    verification, verification_token, error_code = (
        crud.verify_certificate_email_otp(
            db=db,
            student_id=otp_data.student_id,
            requester_email=otp_data.requester_email,
            otp_code=otp_data.otp_code,
        )
    )

    error_details = {
        "OTP_NOT_FOUND": (
            status.HTTP_400_BAD_REQUEST,
            "OTP was not found. Please request a new OTP.",
        ),
        "OTP_EXPIRED": (
            status.HTTP_400_BAD_REQUEST,
            "OTP has expired. Please request a new OTP.",
        ),
        "OTP_TOO_MANY_ATTEMPTS": (
            status.HTTP_429_TOO_MANY_REQUESTS,
            "Too many incorrect OTP attempts. Please request a new OTP.",
        ),
        "OTP_INVALID": (
            status.HTTP_400_BAD_REQUEST,
            "OTP is incorrect",
        ),
    }

    if error_code:
        status_code, detail = error_details[error_code]
        raise HTTPException(
            status_code=status_code,
            detail=detail,
        )

    return schemas.CertificateOtpVerifyResponse(
        message="Email verified successfully",
        email_verification_token=verification_token,
        expires_at=verification.expires_at,
    )


@app.post(
    "/api/certificate-requests",
    response_model=schemas.CertificateRequestResponse,
)
def create_certificate_request(
    request_data: schemas.CertificateRequestCreate,
    db: Session = Depends(get_db),
):
    student = crud.get_student(db, request_data.student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    if student.graduation_status != "GRADUATED":
        raise HTTPException(
            status_code=400,
            detail="Student has not graduated yet",
        )

    if not student.blockchain_tx_id:
        raise HTTPException(
            status_code=400,
            detail="Student has not been verified on blockchain yet",
        )

    try:
        fabric_result = verify_student_on_blockchain(student)
    except Exception:
        logger.exception(
            "Blockchain verification failed for certificate request"
        )
        raise HTTPException(
            status_code=503,
            detail="Blockchain verification is unavailable",
        )

    verification_status = fabric_result["verification_status"]
    is_mismatch = verification_status == "Mismatch"

    if student.is_mismatch != is_mismatch:
        student.is_mismatch = is_mismatch
        db.commit()
        db.refresh(student)

    if verification_status != "Verified":
        raise HTTPException(
            status_code=409,
            detail=(
                "Student data is not verified on blockchain. "
                "Cannot create certificate request."
            ),
        )

    email_verification = crud.consume_certificate_email_verification(
        db=db,
        student_id=request_data.student_id,
        requester_email=request_data.requester_email,
        verification_token=request_data.email_verification_token,
        commit=False,
    )

    if not email_verification:
        raise HTTPException(
            status_code=400,
            detail=(
                "Email has not been verified or the verification token "
                "has expired."
            ),
        )

    try:
        certificate_request = crud.create_certificate_request(
            db=db,
            request_data=request_data,
            commit=False,
        )

        db.commit()
        db.refresh(certificate_request)
    except Exception:
        db.rollback()

        logger.exception(
            "Failed to create certificate request for student %s",
            request_data.student_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Unable to create certificate request.",
        )

    logger.info(
        "Certificate request %s created for student %s",
        certificate_request.id,
        certificate_request.student_id,
    )

    try:
        notify_admins_new_certificate_request(
            db=db,
            request_id=certificate_request.id,
            student_id=certificate_request.student_id,
            requester_name=certificate_request.requester_name,
            requester_email=certificate_request.requester_email,
        )
    except Exception:
        logger.exception(
            "Failed to notify admin for certificate request %s",
            certificate_request.id,
        )

    return certificate_request


@app.get(
    "/api/admin/certificate-requests",
    response_model=schemas.CertificateRequestListResponse,
)
def admin_list_certificate_requests(
    status: str | None = None,
    search: str | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    return crud.list_certificate_requests(
        db=db,
        status=status,
        search=search,
        page=page,
        page_size=page_size,
    )


REQUIRED_CERTIFICATE_PRIVATE_FIELDS = {
    "citizen_id": "Số CCCD/CMND",
}


def ensure_student_can_receive_certificate_document(
    db: Session,
    student: models.Student,
) -> dict:
    try:
        verification = verify_student_on_blockchain(student)
    except Exception:
        logger.exception(
            "Blockchain verification failed before certificate approval/print"
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "BLOCKCHAIN_VERIFICATION_UNAVAILABLE",
                "message": (
                    "Không thể kiểm tra blockchain lúc này. "
                    "Vui lòng thử lại sau."
                ),
            },
        )

    verification_status = verification["verification_status"]
    is_mismatch = verification_status == "Mismatch"

    if student.is_mismatch != is_mismatch:
        student.is_mismatch = is_mismatch
        db.flush()

    if verification_status != "Verified":
        raise HTTPException(
            status_code=409,
            detail={
                "code": "STUDENT_DATA_NOT_VERIFIED",
                "message": (
                    "Không thể duyệt hoặc in giấy xác nhận vì dữ liệu sinh viên "
                    "chưa khớp với blockchain. Vui lòng kiểm tra và đồng bộ lại "
                    "dữ liệu sinh viên trước."
                ),
            },
        )

    private_data = crud.get_student_private_data(student)
    missing_fields = [
        field
        for field in REQUIRED_CERTIFICATE_PRIVATE_FIELDS
        if not private_data.get(field)
    ]

    if missing_fields:
        missing_labels = [
            REQUIRED_CERTIFICATE_PRIVATE_FIELDS[field]
            for field in missing_fields
        ]

        raise HTTPException(
            status_code=409,
            detail={
                "code": "PRIVATE_DATA_INCOMPLETE",
                "message": (
                    "Không thể duyệt hoặc in giấy xác nhận vì sinh viên còn "
                    "thiếu thông tin định danh: "
                    f"{', '.join(missing_labels)}. "
                    "Vui lòng cập nhật số CCCD/CMND trước."
                ),
                "missing_fields": missing_fields,
            },
        )

    return private_data


@app.post(
    "/api/admin/certificate-requests/{request_id}/approve",
    response_model=schemas.CertificateRequestResponse,
)
def admin_approve_certificate_request(
    request_id: int,
    approve_data: schemas.CertificateRequestApprove,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    certificate_request = crud.get_certificate_request(db, request_id)

    if not certificate_request:
        raise HTTPException(
            status_code=404,
            detail="Certificate request not found",
        )

    if certificate_request.status != "CERTIFICATE_PENDING":
        raise HTTPException(
            status_code=400,
            detail="Certificate request has already been reviewed",
        )

    student = crud.get_student(db, certificate_request.student_id)

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    ensure_student_can_receive_certificate_document(db, student)

    updated_request = crud.approve_certificate_request(
        db=db,
        certificate_request=certificate_request,
        admin_note=approve_data.admin_note,
        reviewed_by=current_admin.username,
        commit=False,
    )

    crud.create_audit_log(
        db=db,
        username=current_admin.username,
        action="APPROVE_CERTIFICATE_REQUEST",
        entity_type="CertificateRequest",
        entity_id=updated_request.id,
        details=(
            "Đã duyệt yêu cầu cấp giấy xác nhận cho sinh viên "
            f"{updated_request.student_id}."
        ),
        commit=False,
    )

    try:
        db.commit()
        db.refresh(updated_request)
    except Exception:
        db.rollback()

        logger.exception(
            "Failed to approve certificate request %s",
            request_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Unable to save certificate approval.",
        )

    logger.info(
        "Certificate request %s approved by %s",
        request_id,
        current_admin.username,
    )

    try:
        notify_requester_certificate_approved(
            db=db,
            requester_email=certificate_request.requester_email,
            student_id=certificate_request.student_id,
            admin_note=certificate_request.admin_note,
        )
    except Exception:
        logger.exception(
            "Failed to send approval email for certificate request %s",
            request_id,
        )
    return updated_request

@app.get(
    "/api/admin/certificate-requests/{request_id}/print-data",
    response_model=schemas.CertificatePrintDataResponse,
)
def admin_get_certificate_print_data(
    request_id: int,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    certificate_request = crud.get_certificate_request(db, request_id)

    if not certificate_request:
        raise HTTPException(
            status_code=404,
            detail="Certificate request not found",
        )

    if certificate_request.status != "CERTIFICATE_APPROVED":
        raise HTTPException(
            status_code=400,
            detail="Certificate request has not been approved",
        )

    student = crud.get_student(
        db,
        certificate_request.student_id,
    )

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found",
        )

    private_data = ensure_student_can_receive_certificate_document(
        db,
        student,
    )

    try:
        crud.create_audit_log(
            db=db,
            username=current_admin.username,
            action="VIEW_CERTIFICATE_PRINT_DATA",
            entity_type="CertificateRequest",
            entity_id=certificate_request.id,
            details=(
                "Đã xem dữ liệu in giấy xác nhận của sinh viên "
                f"{student.student_id}."
            ),
        )
    except Exception:
        logger.exception(
            "Failed to create audit log for certificate request %s",
            request_id,
        )
        raise HTTPException(
            status_code=500,
            detail="Unable to record document access.",
        )

    return schemas.CertificatePrintDataResponse(
        request_id=certificate_request.id,
        student_id=student.student_id,
        full_name=student.full_name,
        date_of_birth=student.date_of_birth,
        citizen_id=private_data["citizen_id"],
        citizen_id_issue_date=private_data.get("citizen_id_issue_date"),
        citizen_id_issue_place=private_data.get("citizen_id_issue_place"),
        permanent_address=private_data.get("permanent_address"),
        institution_name=student.institution_name,
        faculty_name=student.faculty_name,
        major=student.major,
        training_mode=student.training_mode,
        entrance_year=student.entrance_year,
        expected_graduation_date=student.expected_graduation_date,
        academic_status=student.academic_status,
        graduation_status=student.graduation_status,
        graduation_date=student.graduation_date,
        graduation_year=student.graduation_year,
        classification=student.classification,
        gpa=student.gpa,
        requested_by=certificate_request.requester_name,
        requested_email=certificate_request.requester_email,
        reason=certificate_request.reason,
        approved_by=(
            certificate_request.reviewed_by
            or current_admin.username
        ),
        approved_at=certificate_request.reviewed_at,
    )

@app.post(
    "/api/admin/certificate-requests/{request_id}/reject",
    response_model=schemas.CertificateRequestResponse,
)
def admin_reject_certificate_request(
    request_id: int,
    reject_data: schemas.CertificateRequestReject,
    db: Session = Depends(get_db),
    current_admin: models.User = Depends(get_current_admin),
):
    certificate_request = crud.get_certificate_request(db, request_id)

    if not certificate_request:
        raise HTTPException(
            status_code=404,
            detail="Certificate request not found",
        )

    if certificate_request.status != "CERTIFICATE_PENDING":
        raise HTTPException(
            status_code=400,
            detail="Certificate request has already been reviewed",
        )

    updated_request = crud.reject_certificate_request(
        db=db,
        certificate_request=certificate_request,
        reject_reason=reject_data.reject_reason,
        reviewed_by=current_admin.username,
        commit=False,
    )

    crud.create_audit_log(
        db=db,
        username=current_admin.username,
        action="REJECT_CERTIFICATE_REQUEST",
        entity_type="CertificateRequest",
        entity_id=updated_request.id,
        details=(
            "Đã từ chối yêu cầu cấp giấy xác nhận cho sinh viên "
            f"{updated_request.student_id}."
        ),
        commit=False,
    )

    try:
        db.commit()
        db.refresh(updated_request)
    except Exception:
        db.rollback()

        logger.exception(
            "Failed to reject certificate request %s",
            request_id,
        )

        raise HTTPException(
            status_code=500,
            detail="Unable to save certificate rejection.",
        )

    logger.info(
        "Certificate request %s rejected by %s",
        request_id,
        current_admin.username,
    )

    try:
        notify_requester_certificate_rejected(
            db=db,
            requester_email=certificate_request.requester_email,
            student_id=certificate_request.student_id,
            reject_reason=certificate_request.reject_reason,
        )
    except Exception:
        logger.exception(
            "Failed to send rejection email for certificate request %s",
            request_id,
        )

    return updated_request


# FABRIC_EXPLORER_V1: public, read-only Fabric Explorer endpoints.
@app.get(
    "/api/explorer/blocks",
    response_model=schemas.ExplorerBlockListResponse,
)
def explorer_blocks(limit: int = Query(default=6, ge=1, le=12)):
    try:
        return list_recent_blocks(limit)
    except Exception:
        logger.exception("Unable to load recent Fabric blocks")
        raise HTTPException(
            status_code=503,
            detail="Blockchain Explorer is temporarily unavailable.",
        )


@app.get(
    "/api/explorer/transactions",
    response_model=schemas.ExplorerTransactionListResponse,
)
def explorer_transactions(limit: int = Query(default=8, ge=1, le=20)):
    try:
        return list_recent_transactions(limit)
    except Exception:
        logger.exception("Unable to load recent Fabric transactions")
        raise HTTPException(
            status_code=503,
            detail="Blockchain Explorer is temporarily unavailable.",
        )


# FABRIC_EXPLORER_NAVIGATION_V1: read-only Explorer detail endpoints.
# FABRIC_EXPLORER_BLOCKS_PAGE_V1: paginated public block list.
@app.get(
    "/api/explorer/blocks/page",
    response_model=schemas.ExplorerBlockPageResponse,
)
def explorer_blocks_page(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=25),
):
    try:
        return list_blocks_page(page, page_size)
    except Exception:
        logger.exception("Unable to load paginated Fabric blocks")
        raise HTTPException(
            status_code=503,
            detail="Blockchain Explorer is temporarily unavailable.",
        )


@app.get(
    "/api/explorer/blocks/{block_number}",
    response_model=schemas.ExplorerBlockResponse,
)
def explorer_block_detail(block_number: int):
    if block_number < 0:
        raise HTTPException(status_code=400, detail="Block number must not be negative.")
    try:
        return get_block_detail(block_number)
    except Exception as exc:
        logger.exception("Unable to load Fabric block detail")
        if "does not exist" in str(exc).lower():
            raise HTTPException(status_code=404, detail="Block does not exist on this channel.")
        raise HTTPException(status_code=503, detail="Blockchain Explorer is temporarily unavailable.")


# FABRIC_EXPLORER_TRANSACTIONS_PAGE_V1: paginated public Graduation transaction list.
@app.get(
    "/api/explorer/transactions/page",
    response_model=schemas.ExplorerTransactionPageResponse,
)
def explorer_transactions_page(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=10, ge=1, le=25),
):
    try:
        return list_transactions_page(page, page_size)
    except Exception:
        logger.exception("Unable to load paginated Fabric transactions")
        raise HTTPException(
            status_code=503,
            detail="Blockchain Explorer is temporarily unavailable.",
        )


@app.get(
    "/api/explorer/transactions/{transaction_id}",
    response_model=schemas.ExplorerTransactionResponse,
)
def explorer_transaction_detail(transaction_id: str):
    try:
        transaction = get_transaction_detail(transaction_id)
    except Exception as exc:
        logger.exception("Unable to load Fabric transaction detail")
        if "format" in str(exc).lower():
            raise HTTPException(status_code=400, detail="Invalid Fabric transaction ID format.")
        raise HTTPException(status_code=503, detail="Blockchain Explorer is temporarily unavailable.")

    if transaction is None:
        raise HTTPException(status_code=404, detail="Transaction was not found in scanned Fabric blocks.")
    return transaction


# EXPLORER_CURRENT_HASH_WARNING_V1
def _explorer_public_current_record(student):
    """Return only the public academic fields already exposed by Explorer."""
    def serialize(value):
        if value is None:
            return None
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return value

    gpa = getattr(student, "gpa", None)
    return {
        "student_id": student.student_id,
        "full_name": student.full_name,
        "institution_code": getattr(student, "institution_code", None),
        "institution_name": student.institution_name,
        "faculty_name": student.faculty_name,
        "major": student.major,
        "training_mode": student.training_mode,
        "degree_id": student.degree_id,
        "degree_type": student.degree_type,
        "entrance_year": student.entrance_year,
        "graduation_status": student.graduation_status,
        "graduation_date": serialize(student.graduation_date),
        "graduation_year": student.graduation_year,
        "classification": student.classification,
        "gpa": str(gpa) if gpa is not None else None,
        "total_credits": getattr(student, "total_credits", None),
        "metadata_hash": student.metadata_hash,
        "metadata_hash_version": student.metadata_hash_version,
        "issuer_msp": None,
        "transaction_id": student.blockchain_tx_id,
        "created_at": serialize(getattr(student, "created_at", None)),
        "updated_at": serialize(getattr(student, "updated_at", None)),
    }


@app.get(
    "/api/explorer/search/{lookup}",
    response_model=schemas.ExplorerSearchResponse,
)
def explorer_search(
    lookup: str,
    db: Session = Depends(get_db),
):
    normalized_lookup = str(lookup or "").strip().lower()

    try:
        result = search_explorer(normalized_lookup)

        if result is None:
            student = crud.get_changed_student_by_metadata_hash(
                db,
                normalized_lookup,
            )

            if student:
                verification = verify_student_on_blockchain(student)

                if verification.get("verification_status") == "Mismatch":
                    transaction_result = (
                        search_explorer(student.blockchain_tx_id)
                        if student.blockchain_tx_id
                        else None
                    )
                    current_hash = (
                        verification.get("metadata_hash_mysql")
                        or student.metadata_hash
                    )
                    confirmed_hash = verification.get(
                        "metadata_hash_blockchain"
                    )
                    result = {
                        "lookup": normalized_lookup,
                        "lookup_type": "current_metadata_hash",
                        "record": _explorer_public_current_record(student),
                        "transaction": (
                            transaction_result.get("transaction")
                            if transaction_result
                            else None
                        ),
                        "verification_status": "Changed",
                        "message": (
                            "Một hoặc nhiều thông tin của người học hiện không giống "
                            "với thông tin đã được nhà trường xác nhận trước đó. "
                            "Kết quả này chưa được xác nhận lại. Vui lòng liên hệ "
                            "nhà trường để kiểm tra."
                        ),
                        "current_metadata_hash": current_hash,
                        "confirmed_metadata_hash": confirmed_hash,
                    }
    except Exception as exc:
        logger.exception("Explorer lookup failed")
        detail = str(exc)
        if "format" in detail.lower() or "64-character" in detail:
            raise HTTPException(status_code=400, detail=detail)
        raise HTTPException(
            status_code=503,
            detail="Hệ thống tra cứu đang tạm thời gián đoạn. Vui lòng thử lại sau.",
        )

    if result is None:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy giao dịch hoặc thông tin người học phù hợp với mã này.",
        )

    return result


@app.get("/api/verify/{student_id}", response_model=schemas.VerifyResponse)
def verify_student(
    student_id: str,
    db: Session = Depends(get_db),
):
    student = crud.get_student(db, student_id)

    if not student:
        # PUBLIC_STUDENT_NOT_FOUND_HTTP_404_V1
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "Không tìm thấy mã sinh viên này trong hệ thống quản lý "
                "của nhà trường."
            ),
        )

    if student.approval_status in {
        "PENDING_BLOCKCHAIN",
        "REJECTED_BLOCKCHAIN",
    }:
        rejected = student.approval_status == "REJECTED_BLOCKCHAIN"
        return schemas.VerifyResponse(
            student_id=student.student_id,
            full_name=None,
            graduation_status="UNKNOWN",
            verification_status=("Not Approved" if rejected else "Pending Approval"),
            message=(
                "Quyết định từ chối đã được ghi nhận trên Blockchain."
                if rejected
                else "Hồ sơ đang chờ quản trị viên duyệt ghi Blockchain."
            ),
            is_graduated=False,
            verified_at=now_utc(),
        )

    if not student.blockchain_tx_id:
        return schemas.VerifyResponse(
            student_id=student.student_id,
            full_name=student.full_name,
            institution_name=student.institution_name,
            faculty_name=student.faculty_name,
            major=student.major,
            training_mode=student.training_mode,
            degree_id=student.degree_id,
            degree_type=student.degree_type,
            entrance_year=student.entrance_year,
            graduation_status=student.graduation_status,
            graduation_date=student.graduation_date,
            graduation_year=student.graduation_year,
            classification=student.classification,
            gpa=student.gpa,
            is_graduated=student.graduation_status == "GRADUATED",
            verification_status="Not Synced",
            message="Student exists in MySQL but has not been synced to blockchain yet.",
            metadata_hash_mysql=student.metadata_hash,
            metadata_hash_blockchain=None,
            metadata_hash_version=student.metadata_hash_version,
            blockchain_tx_id=None,
            verified_at=now_utc(),
        )

    try:
        result = verify_student_on_blockchain(student)
    except Exception as exc:
        logger.exception(
            "Blockchain verify failed for student %s",
            student_id,
        )
        raise HTTPException(
            status_code=503,
            detail=f"Blockchain error: {str(exc)}",
        )

    verification_status = result["verification_status"]

    student.is_mismatch = verification_status == "Mismatch"
    db.commit()
    db.refresh(student)

    return schemas.VerifyResponse(
        student_id=student.student_id,
        full_name=student.full_name,
        institution_name=student.institution_name,
        faculty_name=student.faculty_name,
        major=student.major,
        training_mode=student.training_mode,
        degree_id=student.degree_id,
        degree_type=student.degree_type,
        entrance_year=student.entrance_year,
        graduation_status=student.graduation_status,
        graduation_date=student.graduation_date,
        graduation_year=student.graduation_year,
        classification=student.classification,
        gpa=student.gpa,
        is_graduated=student.graduation_status == "GRADUATED",
        verification_status=verification_status,
        message=result["message"],
        metadata_hash_mysql=result.get("metadata_hash_mysql", student.metadata_hash),
        metadata_hash_blockchain=result.get("metadata_hash_blockchain"),
        metadata_hash_version=student.metadata_hash_version,
        blockchain_tx_id=student.blockchain_tx_id,
        verified_at=now_utc(),
    )
